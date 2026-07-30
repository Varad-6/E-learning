import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

def update_tracker():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Production Readiness Sprint"

    # Headers
    headers = ["Phase", "Task / Feature", "Component Scope", "Status", "Verification Details"]
    ws.append(headers)

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    rows = [
        ["Phase 0", "Diagnostic Sweep & Root Cause Analysis", "Full-stack Audit", "COMPLETED", "Identified ADMIN role scoping bug, MCQ/MSQ schema gap, and missing AI department seed."],
        ["Phase 1", "Core Flow Fix: Department creation -> user section cascade", "Backend / API", "COMPLETED", "Updated RequireRoles & is_global_admin across dependencies, reporting, user_management."],
        ["Phase 1", "Core Flow Fix: Courses page catalog rendering", "Frontend / Backend", "COMPLETED", "Mapped 'published' status to approved catalog items in ViewCourses.tsx & course_service."],
        ["Phase 1", "Core Flow Fix: Reporting page visibility for Admin", "Backend / API", "COMPLETED", "Updated is_global_admin in reporting.py to recognize ADMIN role."],
        ["Phase 1", "Core Flow Fix: MCQ & MSQ question options & auto-grading", "Full-stack", "COMPLETED", "Added options & correct_answer JSON columns, radio/checkbox UI, and auto-scoring."],
        ["Phase 2", "Database Hardening: FK indexes, constraints & pooling", "Database / Alembic", "COMPLETED", "Applied migration f20000000001 with 11 B-tree FK indexes and pool_size=20."],
        ["Phase 3", "Email OTP Integration: Hashed OTP, transactional email, 60s rate limit", "Auth / Services", "COMPLETED", "PasswordResetOTP table, SMTP email service with fallback, 60s rate limit, and OTP UI."],
        ["Phase 4", "Mandatory Courses: Auto-enrollment, visual badge, compliance widget", "Full-stack", "COMPLETED", "is_mandatory column, Creator Studio toggle, auto-enrollment, badge, and compliance card."],
        ["Phase 5", "AI PDF Exam Generation: pypdf parsing, LLM/NLP prompt, staging review", "AI / Backend / UI", "COMPLETED", "AIExamService parsing, pending_ai_questions staging table, upload button & review in ExamCreator."],
        ["Phase 6", "Employee Leaderboard Tab: Department scoping & self-rank highlighting", "Frontend / API", "COMPLETED", "Enabled Leaderboard link in Navbar for Employee, department scoping, (YOU) badge."],
        ["Phase 7", "Badge System Verification: Milestone triggers (courses & exams)", "Backend / Services", "COMPLETED", "BadgeService milestone evaluation across course completions and graded exams."],
        ["Phase 8", "Seed AI Department Test Data: AI dept, 5 courses, 5 exams, 5 users", "Database / Seed", "COMPLETED", "Created seed_ai_department.py and seeded complete AI dataset with attempt history."],
        ["Phase 9", "Console & Network Cleanup", "Full-stack Cleanup", "COMPLETED", "Removed unused difficulty-distribution api call, fixed NameError in exams API for Employee/Manager, and cleaned up console logs."]
    ]

    pass_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    pass_font = Font(name="Calibri", size=10, bold=True, color="166534")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for row_idx, row_data in enumerate(rows, 2):
        ws.append(row_data)
        for col_idx in range(1, 6):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if col_idx in [1, 4]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
            
            if col_idx == 4 and cell.value == "COMPLETED":
                cell.fill = pass_fill
                cell.font = pass_font

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 50
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 65

    wb.save("tracker.xlsx")
    print("✅ tracker.xlsx updated successfully.")

if __name__ == "__main__":
    update_tracker()
